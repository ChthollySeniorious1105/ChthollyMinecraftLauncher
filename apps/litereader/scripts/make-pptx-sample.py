from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
import sys

out = sys.argv[1]
img = sys.argv[2]
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 1 标题页
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'LiteReader 演示文稿样例'
s.placeholders[1].text = 'PPTX 渲染测试 · 2026'
s.notes_slide.notes_text_frame.text = '这是第一张幻灯片的演讲者备注。'

# 2 项目符号
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '功能列表'
tf = s.placeholders[1].text_frame
tf.text = '母版 / 版式继承'
for t, lvl in [('主题颜色与字体', 0), ('二级项目：渐变填充', 1), ('三级项目', 2), ('表格、图片、形状、图表', 0)]:
    p = tf.add_paragraph(); p.text = t; p.level = lvl
r = tf.paragraphs[0].runs[0]; r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
s.notes_slide.notes_text_frame.text = '讲解继承机制。'

# 3 形状
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '形状与连接线'
x = Inches(0.6)
for i, shp in enumerate([MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL, MSO_SHAPE.ISOSCELES_TRIANGLE, MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.CHEVRON, MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.HEXAGON, MSO_SHAPE.HEART]):
    sh = s.shapes.add_shape(shp, x + Inches(1.55) * i, Inches(2.2), Inches(1.35), Inches(1.35))
    sh.text = str(i + 1)
g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.3), Inches(5), Inches(1.6))
g.fill.gradient(); g.fill.gradient_angle = 0
g.fill.gradient_stops[0].color.rgb = RGBColor(0x25, 0x63, 0xEB); g.fill.gradient_stops[1].color.rgb = RGBColor(0xEC, 0x48, 0x99)
g.text = '渐变填充 + 旋转'; g.rotation = 5
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.5), Inches(4.5), Inches(11), Inches(5.8))
c.line.width = Pt(3); c.line.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
ln = c.line._get_or_add_ln(); te = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle'}); ln.append(te)

# 4 表格 + 图片
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '表格与图片'
rows, cols = 4, 3
t = s.shapes.add_table(rows, cols, Inches(0.6), Inches(1.8), Inches(6.5), Inches(2.4)).table
for ri, row in enumerate([['格式', '引擎', '状态'], ['PPTX', 'JSZip + 自研', '✓'], ['HTML', 'iframe 沙箱', '✓'], ['TeX', 'KaTeX', '✓']]):
    for ci, v in enumerate(row): t.cell(ri, ci).text = v
s.shapes.add_picture(img, Inches(7.6), Inches(1.8), width=Inches(5))
tb = s.shapes.add_textbox(Inches(0.6), Inches(5.2), Inches(12), Inches(1))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = '访问官网 '; r.font.size = Pt(20)
r2 = p.add_run(); r2.text = 'example.com'; r2.hyperlink.address = 'https://example.com'; r2.font.size = Pt(20)

# 5 图表
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '季度数据'
cd = CategoryChartData(); cd.categories = ['Q1', 'Q2', 'Q3', 'Q4']
cd.add_series('2025', (12, 18, 9, 22)); cd.add_series('2026', (15, 21, 14, 27))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(11), Inches(5.4), cd)

# 6 隐藏页 + 背景色
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1.5))
tb.text_frame.text = '深色背景 · 隐藏的幻灯片'
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
s._element.set('show', '0')
prs.save(out)
print('saved', out)
